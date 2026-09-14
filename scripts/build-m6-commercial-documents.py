from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(r"E:\face")
SOURCE = ROOT / "docs" / "commercial-v1.2"
OUTPUT = ROOT / "docs" / "commercial-v1.3"
OUTPUT.mkdir(parents=True, exist_ok=True)

ROSE = "B34F69"
DARK = "23171C"
MUTED = "6E5C64"
LIGHT = "F5EFF1"


def set_run_font(run, size=10.5, bold=False, color=DARK):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_title(doc, title, subtitle):
    doc.add_page_break()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(90)
    run = paragraph.add_run(title)
    set_run_font(run, 24, True, ROSE)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(subtitle)
    set_run_font(run, 12, False, MUTED)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("基线日期：2026-08-02｜状态：M6 已验收 / Go（受生产配置门禁约束）")
    set_run_font(run, 9.5, False, MUTED)
    doc.add_page_break()


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(level=level)
    run = paragraph.add_run(text)
    set_run_font(run, 16 if level == 1 else 12, True, ROSE if level == 1 else DARK)
    return paragraph


def add_body(doc, text, bullet=False):
    paragraph = doc.add_paragraph(style="List Bullet" if bullet else None)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.35
    run = paragraph.add_run(text)
    set_run_font(run)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, ROSE)
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, 9.5, True, "FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if len(table.rows) % 2 == 0:
                shade(cells[index], LIGHT)
            run = cells[index].paragraphs[0].add_run(str(value))
            set_run_font(run, 9)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = width
    doc.add_paragraph()


def build_docx(source_name, output_name, role, sections):
    destination = OUTPUT / output_name
    shutil.copy2(SOURCE / source_name, destination)
    doc = Document(destination)
    add_title(doc, f"V1.3 增补：M6 商用运营与发布验收", role)
    for kind, payload in sections:
        if kind == "h1":
            add_heading(doc, payload, 1)
        elif kind == "h2":
            add_heading(doc, payload, 2)
        elif kind == "p":
            add_body(doc, payload)
        elif kind == "b":
            add_body(doc, payload, True)
        elif kind == "t":
            add_table(doc, *payload)
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if "V1.2" in paragraph.text:
                for run in paragraph.runs:
                    if "V1.2" in run.text:
                        run.text = run.text.replace("V1.2", "V1.3")
    doc.core_properties.title = output_name.removesuffix(".docx")
    doc.core_properties.subject = "FACE 美容院店铺管理系统 M6 商用升级验收基线"
    doc.core_properties.comments = "V1.2 原文保留；本文件追加 M6 已验收增补。"
    doc.save(destination)
    return destination


common_status = (
    "M6 已按 M6-01 至 M6-09 顺序完成：经营分析、不可变报表快照、营销同意与独立审批、员工培训资格、"
    "只读开放目录、容量、安全和备份恢复均有真实运行证据。旧 v1/v2 接口继续兼容；M6 不回写 M5 不可变流水。"
)

database_sections = [
    ("h1", "1. M6 数据库交付"),
    ("p", common_status),
    ("t", (["迁移", "对象", "关键约束"], [
        ["V2026080101", "analytics_report_snapshot", "内容哈希、创建幂等、本人下载与过期语义"],
        ["V2026080102", "6 张营销治理表", "明确同意、受众冻结、职责分离、投递唯一"],
        ["V2026080103", "6 张培训/开放平台表", "课程版本、培训历史、密钥摘要、Nonce 防重放"],
    ], None)),
    ("h1", "2. 当前数据库基线"),
    ("b", "最新 Flyway：2026080103，SUCCESS=1；全新 MySQL 8 实例迁移成功。"),
    ("b", "当前基线共 101 张表；M6 新增 13 张表、11 项权限，不删除或改名历史字段。"),
    ("b", "集成完整密钥持久化行数为 0；开放目录敏感字段为 0。"),
    ("h1", "3. 备份、恢复与回滚"),
    ("p", "最终隔离演练使用单事务备份，SHA-256 为 0AD8194069D4FA1F501B9431DB833EA245AC3881D058D5F74707EB3F5921221A。源库与恢复库关键汇总一致，恢复后 readiness=UP，RPO=0，RTO=68.73 秒。"),
    ("p", "生产回滚优先关闭新入口并回退应用，保留新增表、审计、不可变流水和历史事件；有业务数据时不得直接执行反向 DDL。"),
]

api_sections = [
    ("h1", "1. M6 已实现接口域"),
    ("p", common_status),
    ("t", (["接口域", "稳定前缀", "边界"], [
        ["经营分析/报表", "/api/v3/analytics", "事实汇总、本人快照、SHA-256 下载审计"],
        ["营销治理", "/api/v3/marketing", "明确同意、冻结受众、独立审批；外部通道未配置不伪成功"],
        ["员工培训", "/api/v3/training", "课程版本、本人提交、非本人验证、证书有效期"],
        ["集成客户端", "/api/v3/integrations/clients", "密钥只显一次、轮换、撤销、绑定门店"],
        ["开放目录", "/api/v3/open/v1/catalog/services", "仅 catalog:read；时间戳、Nonce、限流与审计"],
    ], None)),
    ("h1", "2. 契约与安全"),
    ("b", "关键写操作继续要求 Idempotency-Key；状态更新继续要求 version；冲突返回 409。"),
    ("b", "开放请求要求 Bearer 密钥、X-Integration-Client、X-Integration-Timestamp、X-Integration-Nonce。"),
    ("b", "401/403/404/409/429 分别表达身份、权限、范围、冲突/重放和限流；错误不回传堆栈或敏感数据。"),
    ("h1", "3. 兼容与发布门禁"),
    ("p", "v1/v2 冻结兼容，M6 仅新增 v3 能力。删除或改名字段仍必须进入新版本接口和迁移窗口；生产渠道、证书、密钥和回调未配置时保持明确不可用。"),
]

implementation_sections = [
    ("h1", "1. M6 阶段完成情况"),
    ("t", (["任务", "交付", "状态"], [
        ["M6-01", "品项分析", "完成"],
        ["M6-02", "品牌分析", "完成"],
        ["M6-03", "门店分析", "完成"],
        ["M6-04", "同比/环比经营比较", "完成"],
        ["M6-05", "不可变报表快照与导出", "完成"],
        ["M6-06", "营销同意、审批与投递治理", "完成"],
        ["M6-07", "员工培训与开放平台", "完成"],
        ["M6-08", "容量、安全、灾备与 WPS 文档同步", "完成"],
        ["M6-09", "最终发布验收、恢复复核与源码快照", "完成"],
    ], None)),
    ("h1", "2. 总体验收事实"),
    ("b", "后端 161/161 测试通过，失败 0、错误 0、跳过 0；后端打包成功。"),
    ("b", "管理端 1782 模块、技师/会员端 199 模块生产构建成功；双端白色主题浏览器验收无错误、无失败请求、无横向溢出。"),
    ("b", "最终容量门禁：20 并发、500 混合请求错误 0；P95=18.05 ms，P99=22.52 ms。"),
    ("b", "管理端与技师/会员端生产依赖漏洞 0；Maven 75 个生产依赖经 OSV 查询漏洞 0，Jackson 已固定为 3.1.5；生产源文件嵌入凭据 0、敏感日志命中 0。"),
    ("b", "最终备份恢复 101 张表汇总一致；RPO=0，RTO=68.73 秒，恢复库 readiness=UP。"),
    ("h1", "3. 发布结论与条件"),
    ("p", "M6 本地商用验收结论为 Go。真实生产发布仍须配置生产数据库凭据、支付/通知渠道密钥与证书、TLS/域名、监控告警、加密备份和逐门店灰度；这些外部条件未满足前保持对应能力不可用，不得伪造成功。"),
]

outputs = [
    build_docx("美容数据库V1.2.docx", "美容数据库V1.3.docx", "数据库设计", database_sections),
    build_docx("美容院店铺管理系统 API接口清单V1.2.docx", "美容院店铺管理系统 API接口清单V1.3.docx", "API 接口契约", api_sections),
    build_docx("美容院店铺管理系统 分阶段开发实施与验收V1.2.docx", "美容院店铺管理系统 分阶段开发实施与验收V1.3.docx", "分阶段开发实施与验收", implementation_sections),
]


def build_prd_pdf():
    appendix = OUTPUT / "m6-prd-appendix.pdf"
    pdfmetrics.registerFont(TTFont("Deng", r"C:\Windows\Fonts\Deng.ttf"))
    pdfmetrics.registerFont(TTFont("DengBold", r"C:\Windows\Fonts\Dengb.ttf"))
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleCN", parent=styles["Title"], fontName="DengBold", fontSize=22,
                           leading=30, textColor=colors.HexColor(f"#{ROSE}"), alignment=TA_CENTER)
    subtitle = ParagraphStyle("SubtitleCN", parent=styles["Normal"], fontName="Deng", fontSize=10,
                              leading=16, textColor=colors.HexColor(f"#{MUTED}"), alignment=TA_CENTER)
    heading = ParagraphStyle("HeadingCN", parent=styles["Heading1"], fontName="DengBold", fontSize=14,
                             leading=20, textColor=colors.HexColor(f"#{ROSE}"), spaceBefore=10, spaceAfter=8)
    body = ParagraphStyle("BodyCN", parent=styles["BodyText"], fontName="Deng", fontSize=9.5,
                          leading=16, textColor=colors.HexColor(f"#{DARK}"), alignment=TA_LEFT, spaceAfter=6)
    small = ParagraphStyle("SmallCN", parent=body, fontSize=8.5, leading=13)
    doc = SimpleDocTemplate(str(appendix), pagesize=A4, leftMargin=20*mm, rightMargin=20*mm,
                            topMargin=18*mm, bottomMargin=18*mm, title="美容PRD V1.3 M6 增补")
    story = [Spacer(1, 42*mm), Paragraph("V1.3 增补：M6 商用运营与发布验收", title),
             Spacer(1, 6*mm), Paragraph("美容院店铺管理系统 PRD｜2026-08-02", subtitle),
             Spacer(1, 4*mm), Paragraph("状态：M6 已验收 / Go（受生产配置门禁约束）", subtitle), PageBreak(),
             Paragraph("1. M6 产品范围与完成情况", heading), Paragraph(common_status, body)]
    rows = [["能力", "交付结论"],
            ["经营分析", "品项、品牌、门店、同比/环比和不可变 CSV 报表快照已验收"],
            ["营销治理", "会员明确同意、独立审批、冻结受众、站内投递真实账本已验收"],
            ["员工培训", "课程版本、员工本人提交、独立验证、证书及有效期已验收"],
            ["开放平台", "仅 catalog:read，密钥只显一次，Nonce、防重放、限流、审计已验收"],
            ["性能/安全/灾备", "20 并发容量门禁、依赖/源码扫描、备份恢复均通过"]]
    table = Table(rows, colWidths=[38*mm, 122*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Deng"), ("FONTNAME", (0, 0), (-1, 0), "DengBold"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{ROSE}")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor(f"#{LIGHT}")), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D7C8CE")),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("LEADING", (0, 0), (-1, -1), 13),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [Spacer(1, 3*mm), table, Paragraph("2. 验收指标", heading),
              Paragraph("后端 161/161 测试通过；管理端 1782 模块、技师/会员端 199 模块生产构建成功。", body),
              Paragraph("20 并发、500 次混合请求错误 0，P95=18.05 ms，P99=22.52 ms；生产依赖高危/严重漏洞为 0。", body),
              Paragraph("数据库最新迁移 2026080103，101 张表；最终备份恢复 RPO=0、RTO=68.73 秒，恢复库 readiness=UP。", body),
              Paragraph("3. 发布门禁", heading),
              Paragraph("本地商用验收为 Go。真实生产仍须完成生产凭据、TLS、支付与通知渠道、监控告警、加密备份和逐门店灰度配置；未配置外部通道继续返回明确不可用，不得伪造成功。", body),
              Paragraph("版本说明：V1.2 原文全部保留，本增补构成 V1.3 当前商用基线；冲突时以不破坏历史数据、保留审计/幂等/版本/冲正和最小权限的规则为准。", small)]
    doc.build(story)
    destination = OUTPUT / "美容PRD_V1.3.pdf"
    writer = PdfWriter()
    for page in PdfReader(SOURCE / "美容PRD_V1.2.pdf").pages:
        writer.add_page(page)
    for page in PdfReader(appendix).pages:
        writer.add_page(page)
    with destination.open("wb") as handle:
        writer.write(handle)
    appendix.unlink()
    return destination


outputs.insert(0, build_prd_pdf())
for output in outputs:
    print(f"M6_COMMERCIAL_DOCUMENT={output}")
print("M6_COMMERCIAL_DOCUMENT_BUILD=PASS")
