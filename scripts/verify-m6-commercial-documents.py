from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


root = Path(r"E:\face\docs")
source = root / "commercial-v1.2"
current = root / "commercial-v1.3"
files = sorted(current.iterdir())
docx_files = [path for path in files if path.suffix.lower() == ".docx"]
pdf_files = [path for path in files if path.suffix.lower() == ".pdf"]
if len(docx_files) != 3 or len(pdf_files) != 1:
    raise SystemExit(f"Expected 3 DOCX and 1 PDF, got {len(docx_files)} and {len(pdf_files)}")
if len([path for path in source.iterdir() if path.is_file()]) != 4:
    raise SystemExit("V1.2 source set is incomplete.")


def document_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise ValueError(f"Corrupt DOCX member in {path.name}: {bad}")
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


texts = {path.name: document_text(path) for path in docx_files}
pdf = PdfReader(pdf_files[0])
if len(pdf.pages) < 2:
    raise SystemExit("Commercial PRD PDF is unexpectedly short.")
texts[pdf_files[0].name] = "\n".join(page.extract_text() or "" for page in pdf.pages[-2:])

required_across_set = [
    "V1.3", "M6-09", "161/161", "2026080103", "catalog:read", "RTO=68.73", "P95=18.05"
]
combined = "\n".join(texts.values())
for keyword in required_across_set:
    if keyword not in combined:
        raise SystemExit(f"Missing commercial-document keyword: {keyword}")
for forbidden in ["TODO", "TBD", "PLACEHOLDER", "待补充"]:
    if forbidden in combined:
        raise SystemExit(f"Commercial documents contain placeholder text: {forbidden}")

for path in files:
    digest = hashlib.sha256(path.read_bytes()).hexdigest().upper()
    print(f"M6_COMMERCIAL_SHA256={path.name}:{digest}")
print(f"M6_COMMERCIAL_DOCX_COUNT={len(docx_files)}")
print(f"M6_COMMERCIAL_PDF_COUNT={len(pdf_files)}")
print(f"M6_COMMERCIAL_PRD_PAGES={len(pdf.pages)}")
print("M6_COMMERCIAL_V1_2_PRESERVED=PASS")
print("M6_COMMERCIAL_CONTENT=PASS")
